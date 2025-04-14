import openpyxl
from PIL import Image, ImageDraw, ImageFont, ImageTk
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import zipfile
import os

class CertificateDesigner:
    def __init__(self, root):
        self.root = root
        self.root.title("Certificate Generator")
        self.root.geometry("1100x900")
        
        # Variables
        self.template_path = ""
        self.excel_path = ""
        self.text_elements = []
        self.dragging = None
        self.template_img = None
        self.photo = None
        self.current_step = 1
        self.excel_headers = []
        self.canvas_scale = 1.0
        
        # Create UI
        self.create_widgets()
        self.show_step(1)
        
    def create_widgets(self):
        # Main container
        self.main_frame = tk.Frame(self.root)
        self.main_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        
        # Step 1: File Selection
        self.step1_frame = tk.Frame(self.main_frame)
        
        tk.Label(self.step1_frame, text="Step 1: Select Files", font=('Arial', 14, 'bold')).pack(pady=10)
        
        # Template selection
        tk.Label(self.step1_frame, text="Certificate Template:").pack()
        tk.Button(self.step1_frame, text="Browse...", command=self.load_template, width=20).pack(pady=5)
        self.template_label = tk.Label(self.step1_frame, text="No template selected", fg="gray")
        self.template_label.pack()
        
        # Excel selection
        tk.Label(self.step1_frame, text="Participant Data (Excel):").pack()
        tk.Button(self.step1_frame, text="Browse...", command=self.load_excel, width=20).pack(pady=5)
        self.excel_label = tk.Label(self.step1_frame, text="No Excel file selected", fg="gray")
        self.excel_label.pack()
        
        # Excel headers preview
        self.headers_label = tk.Label(self.step1_frame, text="", wraplength=400)
        self.headers_label.pack(pady=10)
        
        # Navigation buttons
        nav_frame1 = tk.Frame(self.step1_frame)
        nav_frame1.pack(pady=20)
        self.next_btn1 = tk.Button(nav_frame1, text="Next Step →", command=lambda: self.show_step(2), 
                                 state=tk.DISABLED, width=15)
        self.next_btn1.pack()
        
        # Step 2: Position Fields
        self.step2_frame = tk.Frame(self.main_frame)
        
        # TOP PANEL WITH GENERATE BUTTON - ADDED AT THE TOP
        top_panel = tk.Frame(self.step2_frame, bg="#f0f0f0", padx=10, pady=10)
        top_panel.pack(fill=tk.X, side=tk.TOP)
        
        # Previous button
        tk.Button(top_panel, 
                 text="← Previous Step", 
                 command=lambda: self.show_step(1), 
                 width=15).pack(side=tk.LEFT, padx=5)
        
        # Spacer
        tk.Frame(top_panel).pack(side=tk.LEFT, expand=True)
        
        # GENERATE BUTTON - NOW AT THE TOP
        self.generate_btn = tk.Button(
            top_panel,
            text="GENERATE CERTIFICATES",
            command=self.generate_certificates,
            state=tk.DISABLED,
            bg="#4CAF50",
            fg="white",
            font=('Arial', 12, 'bold'),
            padx=20,
            pady=10,
            borderwidth=3,
            relief=tk.RAISED
        )
        self.generate_btn.pack(side=tk.RIGHT, padx=5)
        
        tk.Label(self.step2_frame, text="Step 2: Position Fields", font=('Arial', 14, 'bold')).pack(pady=10)
        
        # Main content frame
        content_frame = tk.Frame(self.step2_frame)
        content_frame.pack(fill=tk.BOTH, expand=True)
        
        # Left panel - controls
        control_panel = tk.Frame(content_frame, width=250, padx=10)
        control_panel.pack(side=tk.LEFT, fill=tk.Y)
        
        # Field Settings
        tk.Label(control_panel, text="Field Settings", font=('Arial', 12)).pack(pady=5, anchor=tk.W)
        
        # Field selection
        field_frame = tk.Frame(control_panel)
        field_frame.pack(fill=tk.X, pady=5)
        tk.Label(field_frame, text="Field:").pack(side=tk.LEFT)
        self.field_var = tk.StringVar()
        self.field_dropdown = ttk.Combobox(field_frame, textvariable=self.field_var, state="readonly", width=18)
        self.field_dropdown.pack(side=tk.LEFT, padx=5)
        
        # Font size
        font_frame = tk.Frame(control_panel)
        font_frame.pack(fill=tk.X, pady=5)
        tk.Label(font_frame, text="Font Size:").pack(side=tk.LEFT)
        self.font_size = tk.Spinbox(font_frame, from_=8, to=120, width=5)
        self.font_size.delete(0, tk.END)
        self.font_size.insert(0, "36")
        self.font_size.pack(side=tk.LEFT, padx=5)
        
        # Color
        color_frame = tk.Frame(control_panel)
        color_frame.pack(fill=tk.X, pady=5)
        tk.Label(color_frame, text="Color:").pack(side=tk.LEFT)
        self.color_var = tk.StringVar(value="black")
        colors = ["black", "red", "blue", "green", "purple"]
        tk.OptionMenu(color_frame, self.color_var, *colors).pack(side=tk.LEFT, padx=5)
        
        # Add Field button
        tk.Button(control_panel, text="Add Field", command=self.add_text_element, width=15).pack(pady=10)
        
        # Current Fields
        tk.Label(control_panel, text="Current Fields", font=('Arial', 12)).pack(pady=5, anchor=tk.W)
        self.fields_listbox = tk.Listbox(control_panel, height=10)
        self.fields_listbox.pack(fill=tk.BOTH, expand=True)
        
        # Remove button
        tk.Button(control_panel, text="Remove Selected", command=self.remove_field, width=15).pack(pady=5)
        
        # Right panel - canvas
        canvas_panel = tk.Frame(content_frame)
        canvas_panel.pack(side=tk.RIGHT, fill=tk.BOTH, expand=True)
        
        # Canvas with scrollbars
        self.h_scroll = tk.Scrollbar(canvas_panel, orient=tk.HORIZONTAL)
        self.h_scroll.pack(side=tk.BOTTOM, fill=tk.X)
        
        self.v_scroll = tk.Scrollbar(canvas_panel)
        self.v_scroll.pack(side=tk.RIGHT, fill=tk.Y)
        
        self.canvas = tk.Canvas(
            canvas_panel, 
            bg='white',
            xscrollcommand=self.h_scroll.set,
            yscrollcommand=self.v_scroll.set
        )
        self.canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        
        self.h_scroll.config(command=self.canvas.xview)
        self.v_scroll.config(command=self.canvas.yview)
        
        # Bind events
        self.canvas.bind("<Button-1>", self.start_drag)
        self.canvas.bind("<B1-Motion>", self.on_drag)
        self.canvas.bind("<ButtonRelease-1>", self.stop_drag)

    def show_step(self, step_num):
        self.current_step = step_num
        self.step1_frame.pack_forget()
        self.step2_frame.pack_forget()
        
        if step_num == 1:
            self.step1_frame.pack(fill=tk.BOTH, expand=True)
        elif step_num == 2:
            self.step2_frame.pack(fill=tk.BOTH, expand=True)
            if self.template_path:
                self.display_template()
    
    def load_template(self):
        self.template_path = filedialog.askopenfilename(filetypes=[("Image Files", "*.png;*.jpg;*.jpeg")])
        if self.template_path:
            self.template_label.config(text=os.path.basename(self.template_path), fg="green")
            self.check_files_loaded()
            
    def load_excel(self):
        self.excel_path = filedialog.askopenfilename(filetypes=[("Excel Files", "*.xlsx;*.xls")])
        if self.excel_path:
            self.excel_label.config(text=os.path.basename(self.excel_path), fg="green")
            try:
                wb = openpyxl.load_workbook(self.excel_path)
                sheet = wb.active
                self.excel_headers = [str(cell.value) for cell in sheet[1] if cell.value]
                self.headers_label.config(text=f"Detected columns: {', '.join(self.excel_headers)}")
                self.field_dropdown['values'] = self.excel_headers
                if self.excel_headers:
                    self.field_var.set(self.excel_headers[0])
            except Exception as e:
                messagebox.showerror("Error", f"Failed to read Excel file:\n{str(e)}")
            self.check_files_loaded()
            
    def check_files_loaded(self):
        if self.template_path and self.excel_path:
            self.next_btn1.config(state=tk.NORMAL)
        else:
            self.next_btn1.config(state=tk.DISABLED)
    
    def display_template(self):
        try:
            # Open the original template
            original_img = Image.open(self.template_path)
            self.original_size = original_img.size
            
            # Calculate scaling factor for display
            max_display_width, max_display_height = 800, 600
            width_ratio = max_display_width / original_img.width
            height_ratio = max_display_height / original_img.height
            self.canvas_scale = min(width_ratio, height_ratio, 1.0)  # Don't scale up
            
            # Create resized version for display
            display_width = int(original_img.width * self.canvas_scale)
            display_height = int(original_img.height * self.canvas_scale)
            self.template_img = original_img.resize((display_width, display_height), Image.Resampling.LANCZOS)
            
            # Create PhotoImage for canvas
            self.photo = ImageTk.PhotoImage(self.template_img)
            
            # Update canvas
            self.canvas.config(
                scrollregion=(0, 0, display_width, display_height),
                width=display_width,
                height=display_height
            )
            self.canvas.delete("all")
            self.canvas.create_image(0, 0, anchor=tk.NW, image=self.photo)
            
        except Exception as e:
            messagebox.showerror("Error", f"Failed to load template:\n{str(e)}")
        
    def add_text_element(self):
        field = self.field_var.get()
        if not field:
            messagebox.showerror("Error", "Please select a field from the dropdown")
            return
            
        try:
            size = int(self.font_size.get())
        except ValueError:
            messagebox.showerror("Error", "Please enter a valid font size")
            return
            
        color = self.color_var.get()
        
        # Default position at center (in canvas coordinates)
        x = self.template_img.width // 2 if hasattr(self, 'template_img') else 100
        y = self.template_img.height // 2 if hasattr(self, 'template_img') else 100
        
        # Create draggable text element
        text_id = self.canvas.create_text(
            x, y, 
            text=field, 
            fill=color, 
            font=("Arial", size), 
            tags="draggable"
        )
        bbox = self.canvas.bbox(text_id)
        rect_id = self.canvas.create_rectangle(
            bbox, 
            outline="blue" if color == "black" else "black",
            tags="draggable"
        )
        
        # Store both canvas coordinates and original image coordinates
        element = {
            'id': (text_id, rect_id),
            'field': field,
            'font_size': size,
            'color': color,
            'canvas_x': x,
            'canvas_y': y,
            'original_x': int(x / self.canvas_scale),
            'original_y': int(y / self.canvas_scale),
            'original_font_size': int(size / self.canvas_scale)
        }
        
        self.text_elements.append(element)
        self.fields_listbox.insert(tk.END, f"{field} (Size: {size}, Color: {color})")
        self.generate_btn.config(state=tk.NORMAL)
        
    def remove_field(self):
        selection = self.fields_listbox.curselection()
        if not selection:
            messagebox.showerror("Error", "Please select a field to remove")
            return
            
        index = selection[0]
        element = self.text_elements.pop(index)
        for item_id in element['id']:
            self.canvas.delete(item_id)
        self.fields_listbox.delete(index)
        
        if not self.text_elements:
            self.generate_btn.config(state=tk.DISABLED)
        
    def start_drag(self, event):
        items = self.canvas.find_overlapping(event.x-5, event.y-5, event.x+5, event.y+5)
        for item in items:
            for i, element in enumerate(self.text_elements):
                if item in element['id']:
                    self.dragging = (i, element)
                    self.drag_offset_x = element['canvas_x'] - self.canvas.canvasx(event.x)
                    self.drag_offset_y = element['canvas_y'] - self.canvas.canvasy(event.y)
                    return
                    
    def on_drag(self, event):
        if self.dragging:
            i, element = self.dragging
            canvas_x = self.canvas.canvasx(event.x)
            canvas_y = self.canvas.canvasy(event.y)
            
            new_x = canvas_x + self.drag_offset_x
            new_y = canvas_y + self.drag_offset_y
            
            dx = new_x - element['canvas_x']
            dy = new_y - element['canvas_y']
            for item_id in element['id']:
                self.canvas.move(item_id, dx, dy)
            
            # Update positions
            self.text_elements[i]['canvas_x'] = new_x
            self.text_elements[i]['canvas_y'] = new_y
            self.text_elements[i]['original_x'] = int(new_x / self.canvas_scale)
            self.text_elements[i]['original_y'] = int(new_y / self.canvas_scale)
            
    def stop_drag(self, event):
        self.dragging = None
        
    def generate_certificates(self):
        try:
            # Load Excel data
            wb = openpyxl.load_workbook(self.excel_path)
            sheet = wb.active
            
            # Create output directory
            output_dir = "certificates_output"
            os.makedirs(output_dir, exist_ok=True)
            
            # Create ZIP file
            zip_filename = os.path.join(output_dir, "certificates.zip")
            
            with zipfile.ZipFile(zip_filename, 'w', zipfile.ZIP_DEFLATED) as zipf:
                for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=1):
                    # Create certificate from original template
                    img = Image.open(self.template_path)
                    draw = ImageDraw.Draw(img)
                    
                    # Add all text elements
                    for element in self.text_elements:
                        field_name = element['field']
                        try:
                            if field_name in self.excel_headers:
                                col_index = self.excel_headers.index(field_name)
                                text = str(row[col_index]) if row[col_index] is not None else ""
                            else:
                                text = field_name  # Use as static text
                            
                            # Use original coordinates and font size
                            font = ImageFont.truetype("arial.ttf", element['original_font_size'])
                            draw.text(
                                (element['original_x'], element['original_y']),
                                text,
                                fill=element['color'],
                                font=font,
                                anchor="mm"  # Center the text at the position
                            )
                        except Exception as e:
                            print(f"Error processing field {field_name}: {str(e)}")
                            continue
                    
                    # Save to file
                    filename = f"certificate_{row_idx}.png"
                    if self.excel_headers and row[0]:
                        filename = f"{row[0]}_certificate.png"
                    
                    cert_path = os.path.join(output_dir, filename)
                    img.save(cert_path)
                    
                    # Add to ZIP
                    zipf.write(cert_path, filename)
                    os.remove(cert_path)
            
            messagebox.showinfo("Success", 
                f"Certificates generated successfully!\n\n"
                f"Saved to: {os.path.abspath(zip_filename)}\n"
                f"Number of certificates: {sheet.max_row - 1}"
            )
            
        except Exception as e:
            messagebox.showerror("Error", f"Failed to generate certificates:\n{str(e)}")

if __name__ == "__main__":
    root = tk.Tk()
    app = CertificateDesigner(root)
    root.mainloop()